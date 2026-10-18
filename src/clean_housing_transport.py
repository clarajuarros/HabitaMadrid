"""Fase 3: SERPAVI y relaciones GTFS, sin imputación ni asignación espacial."""
from collections import Counter
import hashlib
import json
import math
from pathlib import Path
from zipfile import ZipFile, BadZipFile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.integration_common import IntegrationError, attach_municipal_context, make_issue
from src.municipalities import normalized_name

HOUSING_YEARS = tuple(range(2011, 2025))
HOUSING_TYPES = {"VC": "Vivienda colectiva", "VU": "Vivienda unifamiliar o rural"}
HOUSING_FIELDS = {
    "testigos": "BI_ALVHEPCO_T{tipo}_{yy}",
    "alquiler_m2_mes_mediana": "ALQM2_LV_M_{tipo}_{yy}",
    "alquiler_m2_mes_p25": "ALQM2_LV_25_{tipo}_{yy}",
    "alquiler_m2_mes_p75": "ALQM2_LV_75_{tipo}_{yy}",
    "alquiler_mes_mediana": "ALQTBID12_M_{tipo}_{yy}",
    "alquiler_mes_p25": "ALQTBID12_25_{tipo}_{yy}",
    "alquiler_mes_p75": "ALQTBID12_75_{tipo}_{yy}",
    "superficie_m2_mediana": "SLVM2_M_{tipo}_{yy}",
    "superficie_m2_p25": "SLVM2_25_{tipo}_{yy}",
    "superficie_m2_p75": "SLVM2_75_{tipo}_{yy}",
}
GTFS_SCOPE = "Fuente parcial: autobuses interurbanos CRTM; no representa toda la red de transporte"


def _sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _require_columns(frame, columns, source):
    absent = set(columns) - set(frame.columns)
    if absent:
        raise IntegrationError(f"{source}: faltan columnas {sorted(absent)}")


def _require_key(frame, column, source):
    if frame[column].eq("").any() or frame[column].duplicated().any():
        raise IntegrationError(f"{source}: {column} debe ser una clave única y no vacía")


def _housing_number(value, field, code, year, kind):
    if value is None or value == "":
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise IntegrationError(f"SERPAVI: valor numérico inválido {code}/{year}/{kind}/{field}: {value!r}")
    return value


def _check_housing_values(values, context):
    for prefix in ("alquiler_m2_mes", "alquiler_mes", "superficie_m2"):
        triplet = [values[prefix + suffix] for suffix in ("_p25", "_mediana", "_p75")]
        present = [value is not None for value in triplet]
        if any(present) and not all(present):
            raise IntegrationError(f"SERPAVI: percentiles incompletos {context}/{prefix}")
        if all(present) and not (0 < triplet[0] <= triplet[1] <= triplet[2]):
            raise IntegrationError(f"SERPAVI: percentiles no positivos o desordenados {context}/{prefix}")
    witnesses = values["testigos"]
    if witnesses is not None and (witnesses < 0 or not float(witnesses).is_integer()):
        raise IntegrationError(f"SERPAVI: testigos negativos o no enteros {context}")
    if values["alquiler_mes_mediana"] is not None and (witnesses is None or witnesses <= 0):
        raise IntegrationError(f"SERPAVI: precio sin testigos positivos {context}")


def prepare_housing(master, housing_dir):
    """Lee de nuevo la hoja original; el CSV de una extracción anterior no se usa."""
    housing_dir = Path(housing_dir)
    path = housing_dir / "serpavi_2011_2024.xlsx"
    manifest_path = housing_dir / "serpavi_descarga.json"
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (ValueError, OSError) as exc:
        raise IntegrationError("SERPAVI: no se pudo leer el manifiesto de descarga") from exc
    digest = _sha256(path)
    if manifest.get("sha256") != digest:
        raise IntegrationError("SERPAVI: la huella SHA-256 del XLSX no coincide con el manifiesto")
    if manifest.get("bytes") is not None and manifest["bytes"] != path.stat().st_size:
        raise IntegrationError("SERPAVI: el tamaño del XLSX no coincide con el manifiesto")
    if not isinstance(manifest.get("url"), str) or not manifest["url"].strip():
        raise IntegrationError("SERPAVI: falta URL de procedencia en el manifiesto")
    _require_columns(master, ["codigo_ine", "municipio", "poblacion", "poblacion_anio",
                              "elegible_hasta_50000"], "maestro municipal")
    _require_key(master, "codigo_ine", "maestro municipal")
    master_names = master.set_index("codigo_ine")["municipio"].to_dict()
    records, seen_codes = [], set()
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if "Municipios" not in workbook.sheetnames:
                raise IntegrationError("SERPAVI: falta la hoja Municipios")
            rows = workbook["Municipios"].iter_rows(values_only=True)
            header = list(next(rows))
            if len(header) != len(set(header)):
                raise IntegrationError("SERPAVI: campos repetidos en la cabecera de Municipios")
            expected = {"CPRO", "CUMUN", "NMUN"}
            for year in HOUSING_YEARS:
                for kind in HOUSING_TYPES:
                    expected.update(pattern.format(tipo=kind, yy=str(year)[-2:])
                                    for pattern in HOUSING_FIELDS.values())
            if not expected.issubset(header):
                raise IntegrationError("SERPAVI: faltan campos esperados en Municipios: " +
                                       str(sorted(expected - set(header))))
            positions = {field: idx for idx, field in enumerate(header)}
            for excel_row, row in enumerate(rows, start=2):
                if str(row[positions["CPRO"]]).zfill(2) != "28":
                    continue
                original_code = row[positions["CUMUN"]]
                code = str(original_code).zfill(5)
                if code not in master_names or code in seen_codes:
                    raise IntegrationError(f"SERPAVI: código municipal desconocido o duplicado: {code}")
                seen_codes.add(code)
                source_name = row[positions["NMUN"]]
                if normalized_name(source_name) != normalized_name(master_names[code]):
                    raise IntegrationError(f"SERPAVI: nombre y código no corresponden: {code}/{source_name}")
                for year in HOUSING_YEARS:
                    for kind, label in HOUSING_TYPES.items():
                        field_map = {field: pattern.format(tipo=kind, yy=str(year)[-2:])
                                     for field, pattern in HOUSING_FIELDS.items()}
                        values = {field: _housing_number(row[positions[source_field]], field, code, year, kind)
                                  for field, source_field in field_map.items()}
                        _check_housing_values(values, f"{code}/{year}/{kind}")
                        available = values["alquiler_mes_mediana"] is not None
                        records.append({
                            "codigo_ine": code, "municipio_fuente": source_name,
                            "anio": year, "tipo": kind, "tipologia": label, **values,
                            "dato_disponible": available,
                            "estado_dato": "publicado" if available else "sin_estadistica_publicada",
                            "naturaleza_precio": "Alquiler declarado de vivienda habitual; estadística histórica",
                            "unidad_alquiler_mes": "EUR/mes", "unidad_alquiler_m2_mes": "EUR/m2/mes",
                            "unidad_superficie": "m2", "fuente": "MIVAU_SERPAVI_2011_2024",
                            "archivo_fuente": path.name, "url_fuente": manifest["url"],
                            "hoja_fuente": "Municipios", "fila_fuente_excel": excel_row,
                            **{"campo_fuente_" + field: source_field for field, source_field in field_map.items()},
                            **{"columna_fuente_" + field: get_column_letter(positions[source_field] + 1)
                               for field, source_field in field_map.items()},
                        })
        finally:
            workbook.close()
    except (ValueError, KeyError, StopIteration, BadZipFile) as exc:
        if isinstance(exc, IntegrationError):
            raise
        raise IntegrationError("SERPAVI: no se pudo interpretar la hoja Municipios") from exc
    if seen_codes != set(master_names):
        raise IntegrationError("SERPAVI: cobertura de códigos distinta del maestro: " +
                               str(sorted(set(master_names) - seen_codes)))
    historic = attach_municipal_context(pd.DataFrame(records), master)
    # El contexto demográfico tiene su propio año, no se confunde con el año del alquiler.
    population = master[["codigo_ine", "poblacion", "poblacion_anio"]]
    for column in ("poblacion", "poblacion_anio"):
        if column in historic.columns:
            historic = historic.drop(columns=column)
    historic = historic.merge(population, on="codigo_ine", how="left", validate="many_to_one")
    historic["testigos"] = pd.array(historic["testigos"], dtype="Int64")
    historic = historic.sort_values(["codigo_ine", "anio", "tipo"]).reset_index(drop=True)
    if historic.duplicated(["codigo_ine", "anio", "tipo"]).any() or len(historic) != len(master) * 28:
        raise IntegrationError("SERPAVI: matriz municipio/año/tipología incompleta o repetida")
    latest_year = max(HOUSING_YEARS)
    latest = historic.loc[historic["anio"].eq(latest_year)].copy()
    eligible_latest = latest.loc[latest["elegible_hasta_50000"].fillna(False)].copy()
    availability = eligible_latest.groupby("codigo_ine")["dato_disponible"].any()
    missing = master.loc[master["codigo_ine"].isin(availability.index[~availability])].copy()
    missing["anio_alquiler"] = latest_year
    missing["estado_dato"] = "sin_estadistica_publicada_en_ninguna_tipologia"
    missing["interpretacion"] = "No implica ausencia de viviendas de alquiler"
    coverage = []
    eligible_historic = historic.loc[historic["elegible_hasta_50000"].fillna(False)]
    for year, frame in eligible_historic.groupby("anio", sort=True):
        entry = {"anio": int(year), "municipios_total": len(master),
                 "municipios_elegibles": int(master["elegible_hasta_50000"].sum()),
                 "alguna_tipologia": int(frame.loc[frame["dato_disponible"], "codigo_ine"].nunique())}
        for kind in HOUSING_TYPES:
            entry[kind] = int(frame.loc[frame["tipo"].eq(kind) & frame["dato_disponible"],
                                       "codigo_ine"].nunique())
        coverage.append(entry)
    issues = []
    for row in missing.itertuples(index=False):
        issues.append(make_issue("SERPAVI", "alquiler_no_publicado", registro=row.codigo_ine,
                                 campo="alquiler_mes_mediana", detalle=(
                                     f"Sin precio publicado en {latest_year} en VC ni VU; no se imputa ni se interpreta como cero.")))
    summary = {"archivo_fuente": str(path), "sha256_fuente": digest,
               "url_fuente": manifest["url"], "extraccion": "Hoja Municipios del XLSX original; CSV previo no utilizado",
               "municipios": len(master), "anios": list(HOUSING_YEARS), "tipologias": HOUSING_TYPES,
               "filas_historicas": len(historic), "anio_actual_estadistico": latest_year,
               "filas_actuales": len(latest), "filas_actuales_elegibles": len(eligible_latest),
               "municipios_elegibles_con_precio": int(availability.sum()),
               "municipios_elegibles_sin_precio": len(missing),
               "cobertura_actual": coverage[-1],
               "controles": ["SHA-256 y tamaño verificados", "Códigos y nombres contrastados con el maestro",
                             "Claves municipio/año/tipología únicas", "Percentiles positivos y ordenados",
                             "Testigos enteros no negativos", "Ausencias conservadas sin imputaciones"],
               "limitacion": "Precios históricos declarados; no son ofertas vigentes ni precios actualizados a 2026"}
    return {"vivienda_historico.csv": historic, "vivienda_2024.csv": latest,
            "vivienda_2024_elegibles.csv": eligible_latest,
            "vivienda_cobertura_anual.csv": pd.DataFrame(coverage),
            "vivienda_sin_alquiler_2024.csv": missing}, issues, summary


def _read_gtfs_table(archive, filename, required):
    if filename not in archive.namelist():
        raise IntegrationError("GTFS: falta " + filename)
    with archive.open(filename) as stream:
        frame = pd.read_csv(stream, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    _require_columns(frame, required, "GTFS " + filename)
    frame["fila_fuente"] = range(2, len(frame) + 2)
    return frame


def _gtfs_coordinates(stops, issues):
    for target, source in (("latitud", "stop_lat"), ("longitud", "stop_lon")):
        stops[target + "_original"] = stops[source]
        stops[target] = pd.to_numeric(stops[source], errors="coerce")
    missing = stops["stop_lat"].eq("") | stops["stop_lon"].eq("")
    valid = stops["latitud"].between(-90, 90) & stops["longitud"].between(-180, 180)
    stops["estado_coordenadas"] = "rango_wgs84_valido_pendiente_control_espacial"
    stops.loc[missing, "estado_coordenadas"] = "coordenadas_ausentes_o_incompletas"
    stops.loc[~valid & ~missing, "estado_coordenadas"] = "coordenadas_invalidas"
    for row in stops.loc[~valid | missing].itertuples(index=False):
        issues.append(make_issue("GTFS_CRTM", row.estado_coordenadas, registro=row.stop_id,
                                 campo="stop_lat/stop_lon", valor_original=f"{row.stop_lat}/{row.stop_lon}",
                                 fila_fuente=row.fila_fuente,
                                 detalle="Se conservan los valores originales; coordenadas operativas vacías."))
    stops.loc[~valid | missing, ["latitud", "longitud"]] = float("nan")
    stops["crs_declarado"] = "EPSG:4326"


def prepare_transport(master, raw_dir, chunk_size=200_000):
    """Normaliza el feed parcial y deriva solo pares ruta/parada, no horarios."""
    path = Path(raw_dir) / "crtm_interurbanos_gtfs.zip"
    issues = []
    try:
        with ZipFile(path) as archive:
            stops = _read_gtfs_table(archive, "stops.txt", ["stop_id", "stop_name", "stop_lat", "stop_lon"])
            routes = _read_gtfs_table(archive, "routes.txt", ["route_id", "route_type"])
            trips = _read_gtfs_table(archive, "trips.txt", ["route_id", "trip_id", "service_id"])
            _require_key(stops, "stop_id", "GTFS stops.txt")
            _require_key(routes, "route_id", "GTFS routes.txt")
            _require_key(trips, "trip_id", "GTFS trips.txt")
            route_ids, stop_ids = set(routes["route_id"]), set(stops["stop_id"])
            bad_routes = set(trips["route_id"]) - route_ids
            if bad_routes:
                raise IntegrationError("GTFS: trips.route_id sin ruta: " + str(sorted(bad_routes)[:10]))
            if "parent_station" not in stops:
                stops["parent_station"] = ""
            parents = set(stops.loc[stops["parent_station"].ne(""), "parent_station"])
            if parents - stop_ids:
                raise IntegrationError("GTFS: parent_station sin registro padre: " + str(sorted(parents - stop_ids)[:10]))
            if (stops["parent_station"].ne("") & stops["parent_station"].eq(stops["stop_id"])).any():
                raise IntegrationError("GTFS: parada con referencia a sí misma como parent_station")
            if "location_type" not in stops:
                stops["location_type"] = ""
            if not stops["location_type"].isin(["", "0", "1", "2", "3", "4"]).all():
                raise IntegrationError("GTFS: location_type fuera de los valores esperados")
            stops["tipo_parada_original"] = stops["location_type"]
            stops["tipo_parada"] = stops["location_type"].replace("", "0")
            stops["nombre"] = stops["stop_name"]
            stops["direccion_fuente"] = stops.get("stop_desc", pd.Series("", index=stops.index))
            _gtfs_coordinates(stops, issues)
            # Un nombre que contiene ESTACIÓN no determina el modo de transporte.
            stops["codigo_ine"] = ""
            stops["estado_asignacion_municipal"] = "pendiente_asignacion_espacial"
            stops = attach_municipal_context(stops, master)
            trip_routes = trips.set_index("trip_id")["route_id"].to_dict()
            pairs, n_stop_times = {}, 0
            if "stop_times.txt" not in archive.namelist():
                raise IntegrationError("GTFS: falta stop_times.txt")
            with archive.open("stop_times.txt") as stream:
                chunks = pd.read_csv(stream, dtype=str, keep_default_na=False, encoding="utf-8-sig",
                                     chunksize=chunk_size)
                for chunk in chunks:
                    _require_columns(chunk, ["trip_id", "stop_id", "stop_sequence"], "GTFS stop_times.txt")
                    unmatched_trips = set(chunk["trip_id"]) - set(trip_routes)
                    unmatched_stops = set(chunk["stop_id"]) - stop_ids
                    if unmatched_trips or unmatched_stops:
                        raise IntegrationError("GTFS: stop_times referencia trips/stops inexistentes: " +
                                               str(sorted(unmatched_trips)[:5]) + "/" + str(sorted(unmatched_stops)[:5]))
                    sequences = pd.to_numeric(chunk["stop_sequence"], errors="coerce")
                    if sequences.isna().any() or (sequences < 0).any() or (sequences % 1 != 0).any():
                        raise IntegrationError("GTFS: stop_sequence ausente, negativo o no entero")
                    chunk["route_id"] = chunk["trip_id"].map(trip_routes)
                    chunk["fila_fuente_stop_times_primera"] = range(n_stop_times + 2, n_stop_times + len(chunk) + 2)
                    unique = chunk.drop_duplicates(["route_id", "stop_id"])
                    for row in unique[["route_id", "stop_id", "trip_id", "fila_fuente_stop_times_primera"]].itertuples(index=False):
                        pairs.setdefault((row.route_id, row.stop_id), {
                            "route_id": row.route_id, "stop_id": row.stop_id,
                            "trip_id_testigo": row.trip_id,
                            "fila_fuente_stop_times_primera": row.fila_fuente_stop_times_primera})
                    n_stop_times += len(chunk)
            pair_frame = pd.DataFrame(list(pairs.values()), columns=["route_id", "stop_id", "trip_id_testigo",
                                                                    "fila_fuente_stop_times_primera"])
            pair_frame = pair_frame.sort_values(["route_id", "stop_id"]).reset_index(drop=True)
            feed_versions = []
            if "feed_info.txt" in archive.namelist():
                feed = _read_gtfs_table(archive, "feed_info.txt", [])
                if "feed_version" in feed:
                    feed_versions = sorted(set(feed["feed_version"]) - {""})
            calendar_rows, service_ids = {}, set()
            for name in ("calendar.txt", "calendar_dates.txt"):
                if name in archive.namelist():
                    calendar = _read_gtfs_table(archive, name, ["service_id"])
                    if calendar["service_id"].eq("").any():
                        raise IntegrationError("GTFS: service_id vacío en " + name)
                    calendar_rows[name] = len(calendar)
                    service_ids.update(calendar["service_id"])
            orphan_services = set(trips["service_id"]) - service_ids
            if orphan_services:
                raise IntegrationError("GTFS: trips.service_id sin definición en calendar/calendar_dates: " +
                                       str(sorted(orphan_services)[:10]))
    except (BadZipFile, pd.errors.ParserError, UnicodeError, ValueError) as exc:
        if isinstance(exc, IntegrationError):
            raise
        raise IntegrationError("GTFS: no se pudo interpretar el archivo") from exc
    for table, member in ((stops, "stops.txt"), (routes, "routes.txt"), (pair_frame, "stop_times.txt + trips.txt")):
        table["fuente"] = "GTFS_CRTM_bus_interurbanos"
        table["archivo_fuente"] = path.name
        table["miembro_fuente"] = member
        table["alcance_fuente"] = GTFS_SCOPE
    stops = stops.sort_values("stop_id").reset_index(drop=True)
    routes = routes.sort_values("route_id").reset_index(drop=True)
    issues.append(make_issue("GTFS_CRTM", "cobertura_parcial", detalle=(GTFS_SCOPE +
                             ". Sin asignación municipal en esta fase; una ausencia no demuestra ausencia de transporte.")))
    issues.append(make_issue("GTFS_CRTM", "horarios_no_evaluados", detalle=(
        "Se verifican relaciones estructurales. No se calculan días activos, frecuencias, conexiones ni tiempos de viaje.")))
    summary = {"archivo_fuente": str(path), "sha256_fuente": _sha256(path),
               "paradas_y_estaciones": len(stops), "rutas": len(routes), "viajes_en_archivo": len(trips),
               "filas_stop_times_validadas": n_stop_times, "relaciones_ruta_parada": len(pair_frame),
               "tipos_location_type": {str(k): int(v) for k, v in Counter(stops["tipo_parada"]).items()},
               "tipos_route_type": {str(k): int(v) for k, v in Counter(routes["route_type"]).items()},
               "versiones_feed": feed_versions, "filas_calendarios": calendar_rows,
               "controles": {"claves_unicas": True, "trips_a_routes": True,
                             "stop_times_a_trips_y_stops": True, "parent_station_existente": True,
                             "trips_a_calendar_o_calendar_dates": True},
               "asignacion_municipal": "Pendiente de fase 4; todos los codigo_ine permanecen vacíos",
               "limitacion": GTFS_SCOPE + "; calendario no evaluado ni horarios calculados"}
    return {"transporte_paradas.csv": stops, "transporte_rutas.csv": routes,
            "transporte_rutas_paradas.csv": pair_frame}, issues, summary


def prepare_housing_transport(master: pd.DataFrame, raw_dir: Path, housing_dir: Path):
    housing_tables, housing_issues, housing_summary = prepare_housing(master, housing_dir)
    transport_tables, transport_issues, transport_summary = prepare_transport(master, raw_dir)
    return {**housing_tables, **transport_tables}, housing_issues + transport_issues, {
        "vivienda": housing_summary, "transporte_gtfs": transport_summary}
